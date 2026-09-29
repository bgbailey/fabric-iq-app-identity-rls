"""Build a locally staged, native-editable customer overview and diagram pack."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import shutil
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

SKILL = Path.home() / ".copilot" / "skills" / "fabric-architect"
THEME = Path.home() / ".copilot" / "skills" / "microsoft-deck" / "assets"
sys.path.insert(0, str(THEME))
sys.path.insert(0, str(SKILL))
from ms_theme import (  # noqa: E402
    INK, MS_BLUE, FABRIC_TEAL, SURFACE, SURFACE_WARM, MUTED_LT, LINE,
    ON_DARK, ON_DARK_MUTED, CARD_DARK, TINT_BLUE, TINT_POS, POSITIVE,
    FONT_BODY, FONT_HEAD, FONT_DISPLAY, FONT_MONO, audit,
)
from fabric_architect.air import AIR  # noqa: E402

STAGE = Path(os.environ["LOCALAPPDATA"]) / "scout" / "build" / "iq-isv-customer"
OUTPUT = STAGE / "v1"
OUTPUT.mkdir(parents=True, exist_ok=True)
SOURCES = {
    "S1": ("Fabric IQ MCP: tools, GA status and delegated authentication",
           "https://learn.microsoft.com/en-us/fabric/iq/connectors/fabric-iq-mcp"),
    "S2": ("Execute DAX Queries: roles, customData, privilege and Arrow",
           "https://learn.microsoft.com/en-us/rest/api/power-bi/datasets/execute-dax-queries"),
    "S3": ("Power BI Embedded: application authentication and tokens",
           "https://learn.microsoft.com/en-us/power-bi/developer/embedded/embed-tokens"),
    "S4": ("GenerateToken: EffectiveIdentity and customData",
           "https://learn.microsoft.com/en-us/rest/api/power-bi/embed-token/generate-token"),
    "S5": ("Embedded cloud RLS: USERNAME and effective identity",
           "https://learn.microsoft.com/en-us/power-bi/developer/embedded/cloud-rls"),
    "S6": ("Azure OpenAI Responses API",
           "https://learn.microsoft.com/en-us/azure/foundry/openai/how-to/responses"),
    "S7": ("Fabric Well-Architected security considerations",
           "https://learn.microsoft.com/en-us/azure/well-architected/microsoft-fabric/security"),
    "E1": ("Synthetic Import-model proof, 28 September 2026",
           "Recorded project evidence; source commit ee6df8be48f2fa5cee6e8b15a2aba00f2646e5bd"),
}
ICON = {
    "model": SKILL / "icons" / "fabric" / "png" / "semantic_model_64_item.png",
    "powerbi": SKILL / "icons" / "fabric" / "png" / "power_bi_48_regular.png",
}
SECURITY = AIR.from_file(STAGE / "security_v2.air.json")
PARITY = AIR.from_file(STAGE / "parity_v1.air.json")
MANIFEST = []


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT_BODY, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name or "Text"
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(5)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.text = line
    return shape


def box(slide, x, y, w, h, fill, line=None, radius=False, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name or "Panel"
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    if radius:
        s.adjustments[0] = 0.10
    s._element.spPr.append(OxmlElement("a:effectLst"))
    return s


def node(slide, x, y, w, h, title, subtitle, *, dark=False, icon=None, name=None):
    group = slide.shapes.add_group_shape()
    group.name = name or title
    s = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.10
    s.fill.solid()
    s.fill.fore_color.rgb = CARD_DARK if dark else TINT_BLUE
    s.line.color.rgb = ON_DARK_MUTED if dark else LINE
    s.line.width = Pt(0.8)
    s._element.spPr.append(OxmlElement("a:effectLst"))
    title_y = y + 0.18
    if icon:
        path = ICON[icon]
        if icon == "powerbi":
            plate = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+(w-.60)/2),
                                          Inches(y+.09), Inches(.60), Inches(.60))
            plate.fill.solid()
            plate.fill.fore_color.rgb = SURFACE
            plate.line.fill.background()
            plate._element.spPr.append(OxmlElement("a:effectLst"))
        with Image.open(path) as im:
            ratio = im.width / im.height
        iw, ih = (0.50 * ratio, 0.50) if ratio <= 1 else (0.50, 0.50 / ratio)
        group.shapes.add_picture(str(path), Inches(x + (w-iw)/2), Inches(y+0.14),
                                 width=Inches(iw), height=Inches(ih))
        title_y = y + 0.73
    # Text helpers accept any object exposing shapes.
    text(group, x+0.12, title_y, w-0.24, 0.42, title, 18, SURFACE if dark else INK,
         True, PP_ALIGN.CENTER, name=f"{name or title}-label")
    text(group, x+0.12, title_y+0.47, w-0.24, h-(title_y-y)-0.53, subtitle,
         15, ON_DARK if dark else MUTED_LT, align=PP_ALIGN.CENTER)
    return group


def arrow(slide, a, b, *, start=3, end=1, elbow=False, name="", dark=False):
    con = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT,
                                     0, 0, Inches(1), Inches(1))
    con.name = name
    con.begin_connect(a, start)
    con.end_connect(b, end)
    con.line.color.rgb = ON_DARK_MUTED if dark else MUTED_LT
    con.line.width = Pt(1.6)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    con._element.spPr.get_or_add_ln().append(tail)
    return con


def line(slide, x1, y1, x2, y2, color=LINE, width=1):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                  Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s._element.spPr.append(OxmlElement("a:effectLst"))
    return s


def base(prs, title, *, dark=False, subtitle=None, sources=(), notes="", kind="content", air=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = INK if dark else SURFACE
    text(s, 0.64, 0.46, 12.05, 0.86, title, 30, SURFACE if dark else INK,
         font=FONT_HEAD, name="Slide title")
    if subtitle:
        text(s, 0.66, 1.34, 11.9, 0.60, subtitle, 17, ON_DARK_MUTED if dark else MUTED_LT)
    if sources:
        text(s, 0.66, 7.09, 12, 0.22, "Sources: " + ", ".join(sources) + " | 28 September 2026",
             9, ON_DARK_MUTED if dark else MUTED_LT)
    reference_text = "\n".join(f"{key}: {SOURCES[key][0]}\n{SOURCES[key][1]}" for key in sources)
    s.notes_slide.notes_text_frame.text = (
        notes + "\n\nEvidence status: product/API documentation is distinct from local implementation "
        "and observed live proof. No production support or customer-deployment equivalence is implied.\n"
        "Classification: our own authored slide, grounded in public Microsoft documentation and/or "
        "reviewed synthetic proof. No customer-private or restricted source slide is reproduced. "
        "Apply the appropriate output sensitivity label before external sharing.\n\n" + reference_text
    )
    if air:
        s.notes_slide.notes_text_frame.text += "\n\nArchitecture assumptions:\n" + "\n".join(air.assumptions)
        s.notes_slide.notes_text_frame.text += "\nEvidence gaps:\n" + "\n".join(air.evidence_gaps)
    MANIFEST.append({"deck": prs.core_properties.title, "slide": len(prs.slides), "title": title,
                     "kind": kind, "sources": list(sources), "classification": "our-own/public-source-derived",
                     "architecture": air.title if air else None})
    return s


def callout(s, value, *, dark=False, y=6.35, size=21):
    box(s, 0.64, y-0.08, 12.05, 0.59, CARD_DARK if dark else TINT_POS)
    text(s, 0.84, y+0.02, 11.65, 0.42, value, size, SURFACE if dark else INK, bold=True)


def security_slide(prs):
    s = base(prs, "Your app owns identity. The model enforces access.", dark=True,
             subtitle="Proposed request path | Keep authorization outside the LLM.",
             sources=("S1", "S2", "S3"), kind="architecture", air=SECURITY,
             notes="Walk left to right: the existing application validates its own session and resolves "
             "the user's entitlements. The AI planner separately supplies DAX. The trusted broker combines "
             "the query with server-owned ExternalAppScope and an opaque application-user key. The model "
             "evaluates CUSTOMDATA() within the RLS role and returns only the permitted rows. The broker's "
             "certificate service principal requires workspace Admin to select roles; without the required "
             "role/context that principal may be unrestricted. Therefore the broker is part of the trusted "
             "boundary, not merely a pass-through. Do not claim native IQ supports external-user app-only "
             "authentication. Metadata credentials are shown separately on the next slide. The model return "
             "is structured authorized data; the full flow subsequently explains only those rows.")
    box(s, 0.64, 2.07, 7.10, 3.91, CARD_DARK, ON_DARK_MUTED, name="zone-application")
    box(s, 9.00, 2.07, 3.69, 3.91, CARD_DARK, ON_DARK_MUTED, name="zone-fabric")
    text(s, 0.85, 2.20, 4.1, 0.35, "Your application responsibility", 16, ON_DARK)
    text(s, 9.19, 2.20, 3.28, 0.35, "Fabric / Power BI", 16, ON_DARK)
    app = node(s, 0.94, 4.25, 2.44, 1.42, "Your app", "Existing login\nand entitlements", dark=True, name="node-app")
    planner = node(s, 4.70, 2.62, 2.66, 1.32, "AI query planner", "IQ schema + LLM", dark=True, name="node-planner")
    broker = node(s, 4.70, 4.45, 2.66, 1.22, "Trusted broker", "Fixed role + user key", dark=True, name="node-broker")
    model = node(s, 9.19, 3.61, 2.56, 2.06, "Semantic Model", "RLS-filtered rows", dark=True, icon="model", name="node-model")
    # Attached elbow connectors preserve node movement in PowerPoint.
    arrow(s, app, broker, elbow=True, name="edge-session", dark=True)
    arrow(s, planner, broker, start=2, end=0, name="edge-dax", dark=True)
    arrow(s, broker, model, elbow=True, name="edge-execute", dark=True)
    text(s, 3.40, 4.09, 1.22, 0.64, "Verified\napp context", 14, ON_DARK, align=PP_ALIGN.CENTER)
    text(s, 6.21, 4.05, 1.38, 0.3, "DAX only", 14, ON_DARK)
    text(s, 7.78, 3.97, 1.10, 0.70, "DAX + role\n+ user key", 14, ON_DARK, align=PP_ALIGN.CENTER)
    callout(s, "Only RLS-filtered rows return to the application and its explanation step.", dark=True)
    return s


def identity_slide(prs):
    s = base(prs, "The application user is not the cloud credential.",
             subtitle="Four separate contexts. No external-user Entra onboarding in this proposed design.",
             sources=("S1", "S2", "S3", "S6"), kind="concept",
             notes="The first row is the user's application identity. A future integration receives that "
             "identity from the existing verified session; the current local UI merely simulates it with a "
             "selector. The opaque normalized issuer/subject-derived key maps to model entitlements. "
             "The IQ caller is a separate delegated work/school account: IQ has no app-only authentication "
             "mode. Query execution uses a separate certificate service principal; inference has its own "
             "caller. None of these cloud credentials is the end user's password, session cookie, or "
             "app user represented in Entra. Never send them to the LLM. A production delegated-IQ "
             "credential lifecycle remains a design decision.")
    rows = [
        ("Application user", "Your current login/session", "Opaque user key drives RLS", True),
        ("IQ metadata caller", "Delegated Entra account", "Reads approved model schema", False),
        ("Query executor", "Service principal + certificate", "Submits DAX with role + key", False),
        ("LLM service caller", "Approved inference credential", "Generates DAX; explains rows", False),
    ]
    for i, (a, b, c, emphasis) in enumerate(rows):
        y = 2.18+i*0.86
        box(s, 0.64, y, 12.05, 0.72, TINT_POS if emphasis else (SURFACE_WARM if i % 2 == 1 else TINT_BLUE))
        text(s, 0.85, y+0.17, 2.73, 0.4, a, 19, POSITIVE if emphasis else INK, True)
        text(s, 3.81, y+0.18, 4.0, 0.39, b, 18)
        text(s, 8.02, y+0.18, 4.40, 0.39, c, 18)
    callout(s, "Keep cloud credentials in the backend; keep your users in your application.")
    return s


def parity_slide(prs):
    s = base(prs, "One entitlement decision can serve reports and chat.", dark=True,
             subtitle="Proposed alignment | Same business access intent, two different delivery contracts.",
             sources=("S2", "S3", "S4", "S5"), kind="architecture", air=PARITY,
             notes="This simplified diagram deliberately abstracts the standard Embedded path: backend "
             "Entra authentication and GenerateToken produce an embed token, which the browser uses for "
             "the report. The token conveys EffectiveIdentity for model RLS. The chat path instead uses "
             "our server-controlled roles/customData query broker and returns Arrow rows, then an "
             "explanation. There is no embed token in the chat query request. Both paths should reuse "
             "the entitlement decision and business model, but may require aligning USERNAME versus "
             "CUSTOMDATA expressions. Actual Embedded/chat parity has not been executed. Logical "
             "lanes show responsibility, not VNet or tenant boundaries.")
    auth = node(s, 0.83, 3.10, 2.92, 1.60, "Authorization", "Existing login\n+ current entitlements", dark=True, name="node-auth")
    report = node(s, 5.31, 2.13, 2.67, 1.88, "Power BI", "Embedded reports", dark=True, icon="powerbi", name="node-report")
    broker = node(s, 5.31, 4.68, 2.67, 1.32, "Query broker", "Conversational analytics", dark=True, name="node-broker")
    model = node(s, 9.75, 3.09, 2.69, 1.98, "Semantic Model", "Measures + aligned RLS", dark=True, icon="model", name="node-model")
    arrow(s, auth, report, elbow=True, name="edge-embed", dark=True)
    arrow(s, auth, broker, elbow=True, name="edge-context", dark=True)
    arrow(s, report, model, elbow=True, name="edge-report-query", dark=True)
    arrow(s, broker, model, elbow=True, name="edge-chat-query", dark=True)
    text(s, 3.87, 2.02, 1.40, 0.97, "Embed token\n+ effective\nidentity", 14, ON_DARK, align=PP_ALIGN.CENTER)
    text(s, 3.89, 5.47, 1.30, 0.72, "Trusted role\n+ user key", 14, ON_DARK, align=PP_ALIGN.CENTER)
    text(s, 8.15, 2.15, 1.43, 0.63, "Report\nqueries", 15, ON_DARK, align=PP_ALIGN.CENTER)
    text(s, 8.05, 5.48, 1.65, 0.37, "executeDaxQueries", 13, ON_DARK, align=PP_ALIGN.CENTER)
    callout(s, "Reuse the authorization decision, not the embed token. Parity is the next proof.", dark=True, size=20)
    return s


def new_deck(title):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = title
    prs.core_properties.subject = "ISV application identities, Fabric IQ and semantic-model RLS"
    prs.core_properties.author = "Brendan Bailey | Microsoft"
    prs.core_properties.keywords = "technical proof; synthetic data; proposed architecture; 2026-09-28"
    return prs


def build_overview():
    p = new_deck("Fabric ISV - Keep your identity, add conversational analytics")
    s = base(p, "Keep your login.\nAdd answers.", dark=True, kind="hero",
             notes="Open with the customer objective: add a conversational experience without replacing "
             "the application identity estate or discarding existing Power BI investment. This is a "
             "proposed custom integration anchored in a proven synthetic query/RLS mechanism, not a "
             "claim that native Fabric IQ accepts Power BI embed tokens or external application users. "
             "The decision is whether to complete a bounded technical integration proof.")
    # Hero title is enlarged without adding decorative framework.
    title = s.shapes[0]
    title.top, title.height, title.width = Inches(1.14), Inches(2.18), Inches(7.7)
    for para in title.text_frame.paragraphs:
        para.font.size, para.font.name = Pt(50), FONT_DISPLAY
    text(s, 0.68, 3.66, 7.9, 1.1, "A Fabric path for ISV conversational analytics\nwith application-owned users and model-enforced access.",
         23, ON_DARK)
    box(s, 9.05, 1.65, 3.57, 3.84, CARD_DARK, radius=True)
    node(s, 9.44, 2.03, 2.77, 1.18, "Your application", "Your identity experience", dark=True)
    text(s, 9.39, 3.53, 2.87, 0.5, "Reports + conversation", 22, FABRIC_TEAL, True, PP_ALIGN.CENTER)
    text(s, 9.55, 4.30, 2.54, 0.65, "One business\naccess decision", 20, ON_DARK, align=PP_ALIGN.CENTER)
    text(s, 0.68, 6.51, 10.7, 0.38, "Customer technical overview | Proposed integration | 28 September 2026", 15, ON_DARK_MUTED)

    s = base(p, "Extend the semantic layer, not the identity estate.",
             subtitle="Build on the application and Power BI capabilities you already operate.",
             sources=("S1", "S3", "S5"),
             notes="The existing login can remain the authentication authority for external application "
             "users. Reuse current entitlement decisions, customer/product scope and business measures. "
             "Add the conversational surface and trusted execution broker, not a second customer directory. "
             "Existing RLS expressions may require an explicit identity-contract adaptation; they are not "
             "promised drop-in. No replacement of Embedded reporting is required. This proposal uses the "
             "Power BI semantic-model endpoint of Fabric IQ, not Ontology or Data Agent. An inference "
             "endpoint is needed, but Foundry Agent Service hosting is optional, not a dependency.")
    text(s, 0.72, 2.13, 5.55, 0.55, "Keep", 32, MS_BLUE, True)
    text(s, 0.74, 2.98, 5.28, 2.58, "Your application login\nCustomer / product entitlements\nPower BI reports and measures\nYour customer experience", 25)
    line(s, 6.49, 2.1, 6.49, 5.8)
    text(s, 7.02, 2.13, 5.0, 0.55, "Add", 32, POSITIVE, True)
    text(s, 7.04, 2.98, 5.24, 2.45, "Natural-language questions\nLive semantic context from IQ\nA trusted RLS query broker", 25)
    callout(s, "A new way to consume the model, not a new identity system.")

    security_slide(p)
    identity_slide(p)

    s = base(p, "IQ supplies context. Your broker controls execution.",
             subtitle="Prepared questions still use live schema and real DAX generation.",
             sources=("S1", "S2", "S6"),
             notes="Read this as a simplified process, not evidence every stage has run. The implementation "
             "initializes MCP, sends notifications/initialized, discovers tools/list and calls the fixed "
             "model's GetSemanticModelSchema. The app projects reviewed schema. Azure OpenAI generates "
             "DAX; IQ itself does not have a DAX-generation tool. The application deliberately does not "
             "dispatch IQ ExecuteQuery or ValueSearch. A local custom execute_dax broker injects the "
             "server-owned role and key, executes the newer REST API, decodes Arrow including HTTP-200 "
             "error rowsets, and sends only authorized rows to a separate stateless explanation call. "
             "A tool allowlist is application dispatch control; the delegated OAuth credential is not "
             "metadata-only. No IQ server modification or product plug-in replacement is implied.")
    steps = [("01", "Read schema", "Fabric IQ MCP"), ("02", "Generate DAX", "Your LLM"),
             ("03", "Bind access", "Trusted broker"), ("04", "Enforce RLS", "Semantic model"),
             ("05", "Explain rows", "Your LLM")]
    for i, (n, a, b) in enumerate(steps):
        x = .72 + i*2.44
        text(s, x, 2.46, 1.7, .64, n, 37, MS_BLUE if i < 2 else POSITIVE, font=FONT_DISPLAY)
        text(s, x, 3.29, 2.17, .83, a, 25, bold=True)
        text(s, x, 4.40, 2.18, .75, b, 18, MUTED_LT)
        if i < 4:
            text(s, x+2.08, 2.58, .28, .5, ">", 23, MUTED_LT)
    callout(s, "Native IQ ExecuteQuery and ValueSearch are excluded from this application's path.", size=19)

    s = base(p, "A user key selects entitlements, not a prompt filter.",
             subtitle="Inside the synthetic model: exact customer / product pairs, enforced by RLS.",
             sources=("S2", "E1"), kind="concept",
             notes="The synthetic Import role ExternalAppScope reads CUSTOMDATA(), matches the app-user "
             "key using EXACT and resolves allowed Scope Keys. User Access is a disconnected entitlement "
             "mapping, not a relationship bridge. Scope then filters Activity; User Access has its own "
             "role filter. Cross-customer access is an explicit set of pairs: A/Home and B/Auto do not "
             "imply A/Auto or B/Home. Missing/unmapped keys yield no allowed scopes. The broker supplies "
             "the key; it is not accepted from generated DAX, arbitrary browser claims or model tool "
             "arguments. ALL/REMOVEFILTERS do not erase engine RLS in the observed proof. Production "
             "entitlement synchronization and storage-mode behavior need separate design.")
    text(s, .74, 2.35, 5.1, .70, "app-user-pairs", 29, MS_BLUE, True, font=FONT_MONO)
    text(s, .75, 3.25, 4.45, 1.28, "CUSTOMDATA() resolves\nthe approved scope set.", 25)
    text(s, .75, 5.05, 4.68, .80, "The LLM cannot add an entitlement.", 21, POSITIVE, True)
    rows=[("Customer", "Product", "This user"), ("A", "Home", "Allowed"), ("B", "Auto", "Allowed"),
          ("A", "Auto", "Not granted"), ("B", "Home", "Not granted")]
    for i, row in enumerate(rows):
        y=2.13+i*.71
        box(s, 6.05, y, 6.47, .67, INK if i==0 else TINT_POS if i in (1,2) else SURFACE_WARM)
        for x,w,v in zip((6.25,8.26,10.23),(1.73,1.68,2.01),row):
            text(s,x,y+.15,w,.35,v,18,SURFACE if i==0 else INK,bold=i==0)
    callout(s, "Security is evaluated by the semantic engine, even when the DAX has no user filter.", size=19)

    parity_slide(p)

    s = base(p, "Reuse the trust pattern, not the embed token.",
             subtitle="App-owns-data Embedded versus custom conversation.",
             sources=("S1","S2","S3","S4","S5"),
             notes="Both designs preserve application authentication and use backend-owned effective "
             "authorization. Embedded uses GenerateToken and EffectiveIdentity to scope report access; "
             "the browser consumes the resulting embed token. Our REST broker uses an Entra app-only "
             "bearer plus roles/customData for each query and returns rows, not a report token. The newer "
             "query API requires workspace Admin for a service principal selecting roles; Embedded "
             "token generation typically needs member/admin rights in the relevant workspaces. IQ "
             "metadata additionally needs delegated authentication. Do not imply service principal "
             "profiles, DirectQuery SSO or multi-model parity are validated. Capacity/licensing must be "
             "reviewed for the actual production topology; this slide makes no SKU promise.")
    columns=(.72,4.11,8.41)
    widths=(3.15,4.05,4.02)
    rows=[
        ("", "Power BI Embedded", "Custom conversational path"),
        ("Customer sign-in", "Existing app login", "Existing app login"),
        ("Authorization carrier", "Embed-token effective identity", "Per-query roles + customData"),
        ("API authentication", "Entra token -> GenerateToken", "Entra token -> executeDaxQueries"),
        ("Delivery", "Embedded report experience", "Authorized rows + explanation"),
        ("Additional dependency", "Report and embed-token lifecycle", "Delegated IQ schema + LLM"),
    ]
    for i,row in enumerate(rows):
        y=2.00+i*.665
        box(s,.64,y,12.05,.635,INK if i==0 else SURFACE_WARM if i%2 else TINT_BLUE)
        for x,w,v in zip(columns,widths,row):
            text(s,x+.08,y+.14,w-.15,.44,v,16 if i else 18,SURFACE if i==0 else INK,bold=i==0)
    callout(s,"An embed token does not authenticate Fabric IQ or the query REST API.",size=20)

    s=base(p,"Align the identity contract before claiming shared RLS.",
           subtitle="Proposed cloud-model alignment | Same opaque key; different request envelopes.",
           sources=("S2","S4","S5"),
           notes="The GenerateToken EffectiveIdentity definition supports customData for cloud models "
           "and live Azure Analysis Services models. The query API also accepts customData. The proposed "
           "alignment sends the same opaque key to a CUSTOMDATA-based role, while still supplying "
           "the username/datasets fields required by the Embedded scenario. These fragments are "
           "illustrative and omit report/dataset IDs and credentials. They are not complete executable "
           "requests or evidence Embedded parity has run. If the existing role uses USERNAME or "
           "USERPRINCIPALNAME, preserve the business entitlement decision but adapt or map the expression "
           "explicitly; do not silently pass arbitrary external subjects as effectiveUsername. Confirm "
           "the customer's actual storage mode and supported identity fields before rollout.")
    for x,heading,body in [
        (.72,"Embedded: effective identity",'username: "app-user-A1"\nroles: ["ExternalAppScope"]\ncustomData: "app-user-A1"\ndatasets: [approvedModel]'),
        (7.06,"Query: server-built request",'query: generatedDax\nroles: ["ExternalAppScope"]\ncustomData: "app-user-A1"')]:
        text(s,x,2.10,5.53,.53,heading,22,bold=True)
        box(s,x,2.98,5.56,2.04,INK,radius=True)
        text(s,x+.18,3.22,5.19,1.60,body,16,ON_DARK,font=FONT_MONO)
    text(s,.76,5.44,11.9,.57,"Existing USERNAME() roles may need adaptation to the CUSTOMDATA() contract.",20)
    callout(s,"Reuse the access policy deliberately. Prove equivalent rows through both experiences.",size=19)

    s=base(p,"The query path already passed the engine-RLS proof.",dark=True,
           subtitle="Observed 28 September 2026 | Synthetic Import model | Same unfiltered DAX",
           sources=("E1","S2"),
           notes="This is recorded live evidence, not a new run in this session. On September 28, 2026, "
           "the deterministic broker harness passed 42 acceptance checks using application certificate "
           "authentication, fixed ExternalAppScope and server-resolved customData. A1 total250/count2; "
           "A2 100/2; A3 350/4; B1 700/1; paired scopes1150/3; no-access0/0 for the coalesced summary. "
           "ALL, REMOVEFILTERS, direct-table projection and entitlement enumeration did not expand "
           "scope in this model. A no-access aggregate can return one summary row while activity count "
           "is zero. Do not equate this with live IQ, LLM correctness, actual Embedded parity, "
           "arbitrary storage-mode compatibility or production support. Amounts are synthetic units, "
           "not revenue, dollars, performance figures or customer data.")
    text(s,.77,2.03,3.89,1.30,"42/42",66,FABRIC_TEAL,font=FONT_DISPLAY)
    text(s,.81,3.43,3.65,.95,"Live RLS\nacceptance checks",25,ON_DARK)
    text(s,.81,5.07,3.75,.85,"ALL / REMOVEFILTERS\ndid not expand access.",18,ON_DARK_MUTED)
    vals=[("A / Home",250,2),("A / Auto",100,2),("A / All",350,4),("B / Home",700,1),("A/Home + B/Auto",1150,3),("No access",0,0)]
    text(s,5.08,2.04,2.28,.31,"Application scope",15,ON_DARK_MUTED)
    text(s,10.52,2.04,1.32,.31,"Amount",15,ON_DARK_MUTED,align=PP_ALIGN.RIGHT)
    text(s,11.85,2.04,.83,.31,"Records",13,ON_DARK_MUTED,align=PP_ALIGN.RIGHT)
    for i,(label,total,count) in enumerate(vals):
        y=2.64+i*.49
        text(s,5.07,y,2.28,.30,label,16,ON_DARK)
        if total:
            box(s,7.56,y+.06,3.0*total/1150,.21,FABRIC_TEAL)
        text(s,10.77,y,1.06,.31,f"{total:,}",17,ON_DARK,align=PP_ALIGN.RIGHT)
        text(s,12.05,y,.43,.31,str(count),17,ON_DARK,align=PP_ALIGN.RIGHT)
    callout(s,"Proven: scoped query execution. Not yet proven: the complete IQ conversational experience.",dark=True,size=18)

    s=base(p,"The remaining step is integration, not a new identity estate.",
           subtitle="A working mechanism is a reason to proceed with a bounded proof, not skip it.",
           sources=("S1","S2","E1"),
           notes="The current app implements IQ schema retrieval, generated DAX, broker execution, "
           "explanation and visible execution events, but combined live validation remains pending. "
           "Do not show the disabled UI as proof of execution. Prioritize three genuine integration "
           "questions: can the full chain run reliably under approved delegated IQ credentials; can "
           "Embedded and chat agree on exact user scope with the same entitlement decision; and can "
           "the actual customer model/storage mode support the identity contract? Service-principal "
           "workspace Admin for the query path and endpoint-wide IQ scopes require deliberate "
           "credential isolation. Security hardening, support confirmation, lifecycle and operational "
           "ownership are rollout gates, not reasons to inflate this proof into a complete application.")
    text(s,.74,2.12,4.83,.55,"Already observed",28,POSITIVE,True)
    text(s,.75,3.07,4.8,1.34,"App-only query execution\nModel-enforced user scope\nPaired grants and no-access",23)
    line(s,6.03,2.12,6.03,5.81)
    text(s,6.61,2.12,5.66,.55,"Demonstrate next",28,MS_BLUE,True)
    text(s,6.64,3.07,5.46,2.02,"Live IQ -> generated DAX -> answer\nEmbedded / chat row-level parity\nCustomer-model compatibility\nCredential ownership and support",22)
    callout(s,"IQ MCP is GA. This custom end-to-end design is a technical proof, not a GA product claim.",size=18)

    s=base(p,"Move forward with one model and a measurable decision.",dark=True,
           subtitle="Keep the current portal. Complete a focused Fabric integration spike.",
           sources=("S1","S2","S3"),
           notes="The proposed ask is to nominate an application integration owner and model owner, select "
           "one representative customer model, and agree a small set of paired scenarios. Complete "
           "the live schema/generation/query/explanation demonstration and compare exact rows and "
           "business measures with Embedded. Success means no end-user directory migration, matching "
           "entitlement outcomes, a visible auditable execution path and explicitly resolved deployment "
           "support/credential questions. Do not invent a duration, customer commitment, ROI or rollout "
           "date. This slide proposes a decision; it does not assert customer approval.")
    text(s,.79,2.20,7.83,1.64,"Reuse what works.\nProve only what is new.",38,SURFACE,font=FONT_DISPLAY)
    text(s,.81,4.29,7.22,1.42,"Existing login + entitlement resolver\nOne semantic model\nReports and conversation, side by side",25,ON_DARK)
    box(s,9.10,2.20,3.25,3.52,CARD_DARK,radius=True)
    text(s,9.34,2.48,2.80,.78,"Decision to earn",24,FABRIC_TEAL,True)
    text(s,9.34,3.52,2.77,1.70,"Same user.\nSame allowed rows.\nUseful answers.",25,ON_DARK)
    text(s,.81,6.54,11.78,.42,"Next commitment: application owner + model owner + agreed parity scenarios.",20,ON_DARK)

    s=base(p,"Technical references: query and identity contracts",
           subtitle="Public Microsoft documentation retrieved 28 September 2026.",
           notes="These are clickable primary sources for the claims in the main narrative. The exact "
           "scenario remains subject to product support and customer-model validation. The deck uses "
           "native authored diagrams, not copied source slides. Source IDs appear in slide footers "
           "and full provenance appears in each slide's notes.",
           kind="appendix")
    for i,k in enumerate(("S1","S2","S3","S4","S5")):
        y=2.01+i*.79
        sh=text(s,.75,y,11.8,.43,f"{k}  {SOURCES[k][0]}",20,MS_BLUE,True)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address=SOURCES[k][1]
        text(s,1.31,y+.43,11.1,.26,SOURCES[k][1].replace("https://learn.microsoft.com/en-us/","learn.microsoft.com/.../"),12,MUTED_LT)
    text(s,.79,6.48,11.83,.47,"Embedded analogy does not establish native IQ support for embed tokens or app-only identity.",17)

    s=base(p,"Evidence, scope and source transparency",
           subtitle="A deliberately narrow proposal, with its limits visible.",
           notes="E1: local reviewed stage1-evidence_v2.json reports liveProofPassed=true and42acceptanceChecks; "
           "source commit ee6df8be48f2fa5cee6e8b15a2aba00f2646e5bd. No environmentIDs or raw cloud traces "
           "are included in the customer deck. Upstream skills-for-fabric reference refreshed at "
           "6c11ad58c25992e5d1435ce7cd80d217d5598a31; skills/fabriciq/SKILL.md, "
           "skills/semantic-model-authoring/SKILL.md andcommon/COMMON-CORE.md used. "
           "Local deck index searched: Power BI Embedded Technical Overview s38/s21 and Embedding Power "
           "BI in Your Application s11/s27 offered analogy, not custom broker proof. Source classifications "
           "were unknown. Live GearUp search failed not_authenticated under intended corporate identity; "
           "no restricted source body or artwork is used. Local/live Fabric blog search succeeded but "
           "selected article5190739 concerns broader ontology integrations; it is not proof of this "
           "semantic-model MCP endpoint and is not used for capability status. Official Fabric icon "
           "ZIP provenance is September16; npm freshness unavailable, so latest asset version is not "
           "asserted. All diagram text and topology are authored. New PowerPoint output must be labelled "
           "by the presenter before sharing.",
           sources=("E1","S6","S7"),kind="appendix")
    text(s,.76,2.13,5.69,.51,"Evidence we have",25,POSITIVE,True)
    text(s,.77,2.92,5.60,2.52,"42 live engine-RLS checks\n6 synthetic users / 6 activity rows\nImport mode; fixed query role\nCurrent public API contracts",22)
    text(s,7.06,2.13,5.40,.51,"Not a claim of",25,MS_BLUE,True)
    text(s,7.07,2.92,5.37,2.52,"Production readiness\nLive end-to-end IQ success\nEmbedded / chat parity\nAll-model or storage-mode support",22)
    text(s,.77,6.12,11.7,.68,"Source discovery: local index searched; live GearUp unavailable. No restricted slides reproduced.\nApply the appropriate sensitivity label before sharing this newly authored file.",15,MUTED_LT)
    return p


def census(prs):
    def descendants(shapes):
        for s in shapes:
            yield s
            if s.shape_type == 6:
                yield from descendants(s.shapes)
    return {
        "slides": len(prs.slides),
        "top_level_shapes": sum(len(s.shapes) for s in prs.slides),
        "all_shapes": sum(1 for s in prs.slides for _ in descendants(s.shapes)),
        "titles": [next(sh.text for sh in s.shapes if sh.name == "Slide title") for s in prs.slides],
    }


def check(prs):
    for n,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if sh.width <= 0 or sh.height <= 0:
                # Horizontal/vertical native connectors legitimately have zero height/width.
                if sh.shape_type != 9:
                    raise ValueError((n,sh.name,"empty dimensions"))
            if sh.left < -Inches(.03) or sh.top < -Inches(.03) or sh.left+sh.width > Inches(13.37) or sh.top+sh.height > Inches(7.53):
                raise ValueError((n,sh.name,"off slide"))
        if not s.notes_slide.notes_text_frame.text.strip():
            raise ValueError((n,"missing notes"))
    return audit(prs, appendix_from=13)


def save(prs, name):
    path=OUTPUT/name
    issues=check(prs)
    if issues:
        raise ValueError(("density",issues))
    prs.save(path)
    actual=census(Presentation(path))
    if actual != census(prs):
        raise ValueError("Staged census differs from authored presentation")
    return {"path":str(path),"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"census":actual}


def main():
    overview=build_overview()
    diagrams=new_deck("ISV security diagrams - editable")
    security_slide(diagrams)
    identity_slide(diagrams)
    parity_slide(diagrams)
    outputs=[save(overview,"Fabric_ISV_Identity_and_Analytics_v1.pptx"),
             save(diagrams,"ISV_Security_Diagrams_v1.pptx")]
    assets=[]
    for name,path in ICON.items():
        assets.append({"name":name,"canonical_path":str(path),"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),
                       "terms":"https://learn.microsoft.com/en-us/fabric/fundamentals/icons",
                       "source":"Microsoft official Fabric icon ZIP, retrieved 2026-09-16"})
    shutil.copy2(STAGE/"security_v2.air.json",OUTPUT/"security_v1.air.json")
    shutil.copy2(STAGE/"parity_v1.air.json",OUTPUT/"parity_v1.air.json")
    result={"outputs":outputs,"slides":MANIFEST,"sources":SOURCES,"assets":assets,
            "source_retrieval_date":"2026-09-28",
            "status":"staged; awaiting visual/semantic review and versioned publish",
            "public_release":False,"native_diagram_editability":"Native grouped nodes and attached PowerPoint connectors; interactive check pending",
            "source_gaps":["Live GearUp not authenticated; no source slides collated.",
                           "Fabric icon npm freshness unavailable; official ZIP provenance retained."],
            "upstream_commit":"6c11ad58c25992e5d1435ce7cd80d217d5598a31"}
    (OUTPUT/"build_manifest_v1.json").write_text(json.dumps(result,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"outputs":outputs,"slides":len(MANIFEST)},indent=2))


if __name__=="__main__":
    main()
