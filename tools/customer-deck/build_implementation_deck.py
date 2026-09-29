"""Version 2: explain the implemented teaching demo, not a proposed customer design."""
from copy import deepcopy
import json
from pathlib import Path
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt
import build_customer_deck as b

b.OUTPUT = b.STAGE / "v2"
b.OUTPUT.mkdir(parents=True, exist_ok=True)
b.SECURITY = b.AIR.from_file(b.STAGE / "implementation_v2.air.json")
b.SOURCES["E2"] = ("Live integrated teaching demo, 28 September 2026 Eastern",
                   "See docs/status.rst: live IQ, GPT-5.4, scoped query and explanation; selected synthetic cases.")


def replace(shape, value):
    first = shape.text_frame.paragraphs[0]
    props = deepcopy(first._p.pPr)
    shape.text_frame.text = value
    for p in shape.text_frame.paragraphs:
        if p._p.pPr is not None:
            p._p.remove(p._p.pPr)
        if props is not None:
            p._p.insert(0, deepcopy(props))


def notes(slide, value, keys):
    slide.notes_slide.notes_text_frame.text = (
        value + "\n\nThis documents an educational implementation, not production certification. "
        "The selector is not authentication. Power BI Embedded is comparison context, not an implemented "
        "report feature. Classification: our own, public documentation and synthetic evidence only. "
        "Microsoft artwork retains its original terms. Check the output label before sharing.\n\n" +
        "\n".join(f"{k}: {b.SOURCES[k][0]}\n{b.SOURCES[k][1]}" for k in keys)
    )


def implementation(s):
    replacements = {
        "Your app owns identity. The model enforces access.": "The running demo binds app-user context in QueryBroker.",
        "Proposed request path | Keep authorization outside the LLM.": "Implemented request path | React + .NET + live IQ + GPT-5.4 + Power BI",
        "Your application responsibility": "React / .NET localhost demo",
        "Verified\napp context": "Synthetic\nuser context",
    }
    for sh in s.shapes:
        if sh.has_text_frame and sh.text in replacements:
            replace(sh, replacements[sh.text])
    groups = {sh.name:sh for sh in s.shapes if sh.shape_type == 6}
    for name,title,subtitle in (
        ("node-app","Demo selector","6 app-owned subjects\nNot authentication"),
        ("node-planner","Live query planner","IQ schema + GPT-5.4"),
        ("node-broker","QueryBroker","ExternalAppScope + key"),
    ):
        texts=[sh for sh in groups[name].shapes if sh.has_text_frame and sh.text]
        replace(texts[0],title)
        replace(texts[1],subtitle)
    notes(s,"Actual request: React sends userId/questionId; DemoRunner resolves SyntheticSubjects "
          "into immutable context. LiveDemoPlanner obtains IQ metadata and calls ResponsesDaxGenerator. "
          "QueryBroker receives DAX plus the hidden context, obtains a certificate app-only token and "
          "calls executeDaxQueries with ExternalAppScope/customData. It decodes Arrow and checks error "
          "rowsets. ResponsesExplainer sees the question and current allowed rows, not credentials or "
          "other users. No end-user authentication is implemented: the selector simulates the result of "
          "a future application's authenticated session. This request-only diagram deliberately omits "
          "return arrows. The existing four-component AIR maps the actual code; Entra credential "
          "boundaries are detailed on the next slide.",("S1","S2","S6","E2"))
    return s


def identity(p):
    s=b.identity_slide(p)
    for sh in s.shapes:
        if sh.has_text_frame:
            changes={
                "Four separate contexts. No external-user Entra onboarding in this proposed design.":"Implemented identity contexts | The app subject is not an Entra account.",
                "Your current login/session":"SyntheticSubjects selector",
                "Your app's existing login":"Demo selector",
                "Keep cloud credentials in the backend; keep your users in your application.":"User context is separate from the credentials used to call cloud services.",
            }
            if sh.text in changes: replace(sh,changes[sh.text])
    notes(s,"The sample uses an opaque selected subject, not real application login. IQ: separately "
          "registered public client, interactive delegated sign-in, then silent protected-cache access "
          "with Item.Read.All, Item.Execute.All and Dataset.Read.All. Query: nonexportable Windows "
          "certificate service principal with Admin only in the isolated model workspace. Inference: "
          "subscription-selected Azure CLI credential; token tenant and principal are verified. The "
          "LLM gets no credentials and cannot set roles/customData. The IQ OAuth scopes are endpoint-wide; "
          "the allowlist is application dispatch control, not metadata-only OAuth.",("S1","S2","S6"))
    return s


def process(p):
    s=b.base(p,"Seven real service requests turn one question into an answer.",
             subtitle="Current implementation | One normal-path question, no DAX or answer template.",
             sources=("S1","S2","S6","E2"))
    labels=[("4 calls","Fabric IQ MCP","Initialize, notify,\nlist tools, read schema"),
            ("1 call","GPT-5.4","Generate DAX\nfrom returned schema"),
            ("1 call","Power BI","Execute with fixed\nrole + customData"),
            ("1 call","GPT-5.4","Explain only the\nRLS-filtered rows")]
    for i,(count,title,body) in enumerate(labels):
        x=.73+i*3.09
        b.text(s,x,2.30,2.72,.62,count,36,b.MS_BLUE,font=b.FONT_DISPLAY)
        b.text(s,x,3.26,2.81,.55,title,24,bold=True)
        b.text(s,x,4.17,2.80,1.26,body,20)
    b.callout(s,"IQ ExecuteQuery and ValueSearch are not dispatched. The broker owns business-data execution.",size=18)
    notes(s,"FabricIqSchemaClient sends initialize, notifications/initialized, tools/list and "
          "GetSemanticModelSchema. Additional tool-list pages cost extra calls. The LLM—not IQ—generates "
          "DAX. Generation and explanation use separate store:false Responses calls. Runtime API paths, "
          "identity categories, timestamps, schema/query hashes and model usage appear in the inspector. "
          "Failed requests consume budget and return explicit errors; no static fallback exists. "
          "Seven is the usual service-call count, not a latency or cost guarantee.",("S1","S2","S6","E2"))
    return s


def build():
    p=b.new_deck("Fabric ISV application-identity RLS - implemented educational demo v2")
    s=b.base(p,"Application-owned users.\nModel-enforced access.",dark=True,kind="hero")
    sh=s.shapes[0];sh.top=Inches(1.01);sh.height=Inches(2.1)
    for para in sh.text_frame.paragraphs: para.font.size=Pt(43);para.font.name=b.FONT_DISPLAY
    b.text(s,.76,3.63,10.98,1.26,"A working educational implementation of\nFabric IQ + generated DAX + a trusted RLS query broker.",26,b.ON_DARK)
    b.text(s,.77,5.76,11.64,.65,"React / .NET 8 / GPT-5.4 / Power BI semantic model\nImplementation walkthrough | 28 September 2026",18,b.ON_DARK_MUTED)
    notes(s,"This deck documents the code and selected real executions of the educational sample, "
          "not a proposed reference architecture. The application user selector is a teaching shortcut "
          "and not authentication. The code illustrates the integration point for an ISV's existing "
          "verified user session. No actual Power BI Embedded report frame is included.",("E2","S1","S2"))
    implementation(b.security_slide(p))
    identity(p)
    process(p)

    s=b.base(p,"CUSTOMDATA() resolves the user's exact entitlement pairs.",
             subtitle="Implemented in the ExternalAppScope model role.",sources=("S2","E1","E2"))
    b.text(s,.77,2.27,5.20,.60,"app-user-pairs",30,b.MS_BLUE,True,font=b.FONT_MONO)
    b.text(s,.79,3.26,5.2,1.47,"User Access -> Scope -> Activity\n\nNo LLM-supplied access rules.",23)
    for i,(label,value) in enumerate((("A / Home","Allowed"),("B / Auto","Allowed"),("A / Auto","Not granted"),("B / Home","Not granted"))):
        y=2.16+i*.77
        b.box(s,6.57,y,5.91,.66,b.TINT_POS if i<2 else b.SURFACE_WARM)
        b.text(s,6.83,y+.14,2.47,.37,label,20,bold=True)
        b.text(s,9.68,y+.14,2.50,.37,value,20)
    b.callout(s,"The broker always attaches the role and key; its privileged credential alone is not row-restricted.",size=17)
    notes(s,"User Access is a disconnected entitlement table. Scope's RLS expression uses "
          "CUSTOMDATA and EXACT to match the subject and scope key, then the one-direction relationship "
          "propagates access to Activity. User Access itself has a role filter. The allowed set is exact "
          "customer/product pairs, not a Cartesian product. Role omission could give the workspace-admin "
          "query service principal unrestricted results; the backend is a trusted component.",("S2","E1","E2"))

    s=b.base(p,"The same question returns the selected user's data.",dark=True,
             subtitle="Observed live: IQ schema -> GPT-5.4 DAX -> RLS -> GPT-5.4 explanation.",sources=("E2",))
    for x,label,amount,count in ((.90,"A / Home","250","2 records"),(5.02,"B / Home","700","1 record"),(9.18,"No access","0","0 records")):
        b.text(s,x,2.28,3.1,.46,label,24,b.ON_DARK,True)
        b.text(s,x,3.27,3.08,1.0,amount,59,b.FABRIC_TEAL,font=b.FONT_DISPLAY)
        b.text(s,x,4.65,3.06,.54,count,24,b.ON_DARK)
    b.callout(s,"Different app subjects. The model—not a prompt filter—restricts the returned rows.",dark=True,size=20)
    notes(s,"Actual observed overview outputs were A1 250/2, B1 700/1, no-access0/0. The A1/B1 "
          "runs happened to use identical generated ROW DAX; generally generated DAX can differ. "
          "The independent 42-check harness proves the fixed-same-query RLS cases. Zero activity count "
          "may still be represented by one aggregate row. Values are synthetic units, not customer money "
          "or performance metrics.",("E2","E1"))

    s=b.base(p,"The demo exposes the query, rows and execution path.",
             subtitle="Five prepared questions; generated DAX on every execution.",sources=("E2",))
    rows=[("Overview","A1: 250 across 2 records"),
          ("Customer / product breakdown","Paired user: A/Home 250 + B/Auto 900"),
          ("Underlying records","A1: activity keys 1 and 2"),
          ("Daily amount and count","A3: 140 / 2 records -> 210 / 2 records"),
          ("Explicit B / Home request","Denied for paired user; B1 returns key 5 / 700")]
    for i,(a,v) in enumerate(rows):
        y=2.02+i*.77
        b.box(s,.66,y,12.02,.69,b.SURFACE_WARM if i%2 else b.TINT_BLUE)
        b.text(s,.88,y+.17,4.05,.35,a,18,bold=True)
        b.text(s,5.00,y+.17,7.26,.35,v,18)
    b.callout(s,"Read the generated DAX and model rows before relying on the explanation.",size=20)
    notes(s,"These are selected live observations, not claims about all30combinations. Click a subject, "
          "select a question and explicitlyRun. Selection alone does not call the cloud. Expand "
          "the inspector to show delegated metadata auth, schema projection, generatedDAX, broker "
          "role/customData and explanation provenance. Switching subjects clears old results and "
          "cancels/stale-gates prior events. No conversation history is persisted in the browser.",("E2",))

    s=b.base(p,"Embedded is the analogy—not a feature hidden in this demo.",
             subtitle="Comparison only | No embed-token or report implementation.",
             sources=("S2","S3","S4","S5"))
    rows=[("Concern","Power BI Embedded","This implementation"),
          ("User context","Authenticated app user","SyntheticSubjects selector"),
          ("RLS carrier","EffectiveIdentity in embed token","roles + customData on each query"),
          ("Delivery","Report iframe / Power BI SDK","Arrow rows + LLM explanation"),
          ("Cloud access","GenerateToken with Entra identity","IQ delegated + query certificate SP")]
    for i,row in enumerate(rows):
        y=2.03+i*.77
        b.box(s,.64,y,12.05,.70,b.INK if i==0 else b.SURFACE_WARM)
        for x,w,value in zip((.87,3.55,8.09),(2.47,4.20,4.25),row):
            b.text(s,x,y+.17,w,.41,value,17,b.SURFACE if i==0 else b.INK,bold=i==0)
    b.callout(s,"An existing USERNAME() role is not automatically a drop-in match for this CUSTOMDATA() role.",size=18)
    notes(s,"Embedded app-owns-data supports arbitrary application authentication. Its backend supplies "
          "EffectiveIdentity toGenerateToken. This sample instead demonstrates a fixed CUSTOMDATA-based "
          "role through executeDaxQueries. Existing app login/entitlement logic can be reused when "
          "integrating, but username-versus-customData semantics must align. Embed tokens are not "
          "bearer tokens forIQorRESTqueries. Actual Embedded/report parity is not implemented or proven.",("S2","S3","S4","S5"))

    s=b.base(p,"The repository teaches the trust boundary in code.",
             subtitle="Small components with explicit responsibilities.",sources=("E2",))
    parts=[("DemoCatalog / SyntheticSubjects","Prepared question text; immutable subject context"),
           ("FabricIqSchemaClient","Real MCP lifecycle, metadata tool allowlist, schema projection"),
           ("ResponsesDaxGenerator","Question-specific LLM input; structured DAX output"),
           ("QueryBroker / Arrow parser","Certificate auth, role/key injection, engine rows"),
           ("ResponsesExplainer / React","Current-row explanation and visible execution events")]
    for i,(a,v) in enumerate(parts):
        y=2.06+i*.74
        b.text(s,.81,y,4.61,.49,a,19,b.MS_BLUE,True)
        b.text(s,5.70,y,6.74,.49,v,18)
    b.callout(s,"Start with docs/learning-guide.rst, then follow the quickstart and source links.",size=20)
    notes(s,"The educational repo includes a five-minute guide, self-contained written deployment steps "
          "for your own development tenant, example configs, exactAPIreferences and an Embedded comparison. "
          "Model packaging is offline; deployment is an explicit separate operation. Credentials, local "
          "approval records, tenant/resource IDs and raw service traces are excluded. Authored code is "
          "MIT; Microsoft icons retain original use terms. Publication/visibility remains a separate action.",("E2","S1","S2"))

    s=b.base(p,"Integration taught us where plausible code was wrong.",
             subtitle="Lessons from the real services, preserved in the educational sample.",sources=("S1","E2"))
    items=[("IQ schema is not the citation envelope.","Read schema from content[].text; validate artifact identity."),
           ("Hidden columns may be omitted.","Project returned metadata; never invent absent fields."),
           ("Example values can become accidental filters.","Provide literal hints only for the question that needs them."),
           (".NET UTC timestamps include +00:00.","The browser accepts both valid UTC forms, not just Z.")]
    for i,(title,body) in enumerate(items):
        y=2.06+i*.91
        b.text(s,.84,y,11.59,.41,title,22,bold=True)
        b.text(s,.85,y+.43,11.58,.36,body,18,b.MUTED_LT)
    notes(s,"An initial details query incorrectly used customerB/Home because global prompt hints "
          "included those values for everyquestion. RLS was still applied, but the business answer was "
          "wrong. The corrected requestscopes hints to the actualcustomerBquestion. Daily wording now "
          "explicitlyasks foramountandcount. The SDK/CLI integrationalsorequired subscription-only "
          "selection plus post-token tenant validation, because AzureCLI rejects tenant andsubscription "
          "together. Theseareeducational integration fixes, not a claim of generalDAX correctness.",("E2",))

    s=b.base(p,"Show the working mechanism, and name what it does not do.",dark=True,
             subtitle="Educational implementation | Not a production application or product-support commitment.",
             sources=("E1","E2"))
    b.text(s,.81,2.14,5.55,.50,"Implemented and demonstrated",26,b.FABRIC_TEAL,True)
    b.text(s,.82,3.03,5.38,2.45,"Live IQ metadata and GPT-5.4\nApplication-context RLS queries\nExact rows and explanations\nExecution inspector",23,b.ON_DARK)
    b.text(s,7.06,2.14,5.41,.50,"Not implemented",26,b.ON_DARK,True)
    b.text(s,7.08,3.03,5.27,2.45,"Real application sign-in\nPower BI Embedded report parity\nPublic web or MCP hosting\nProduction lifecycle / scale",23,b.ON_DARK)
    b.callout(s,"The ISV integration point is the trusted subject resolver—not a replacement identity directory.",dark=True,size=18)
    notes(s,"Do not confuse app-owned identity with completed login. The synthetic selector is the "
          "only user entrypoint today. To adapt thissample, replace its resolverinput with a verified "
          "application session and an approved entitlementmapping, not a client-supplied rawkey. "
          "This repo's value is a working, inspectableteachingpath ratherthan a hardenedproduct. "
          "The full30case matrixis intentionallynotrequiredfor this educationaldelivery.",("E1","E2"))

    s=b.base(p,"Source contracts and a reproducible starting point",
             subtitle="Public Microsoft documentation + actual sample code + dated synthetic observations.",
             kind="appendix")
    for i,key in enumerate(("S1","S2","S3","S4","S6")):
        y=2.04+i*.72
        sh=b.text(s,.82,y,11.65,.48,f"{key}  {b.SOURCES[key][0]}",20,b.MS_BLUE,True)
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address=b.SOURCES[key][1]
    b.text(s,.84,6.08,11.62,.74,"Repository guide: docs/quickstart.rst | Evidence: docs/status.rst\nAll slides are authored from public sources and synthetic data. Check the sharing label.",17)
    notes(s,"PublicAPI linksground supportedcontracts,notproductioncertification. LiveGearUpsearch was "
          "retried successfully for thisrevision and localdeckindexdiscoverywas retained. Genericplatform "
          "slides were notcollated because this deckdocuments specificcode; no restriction-based source "
          "selection occurred. Upstreamskills-for-fabriccommit6c11ad58c25992e5d1435ce7cd80d217d5598a31. "
          "Diagram authored/reviewed with scopedGPT6Astra. Official Fabricicons fromcanonical skillcache "
          "retainoriginalterms. Allruntimeevidenceis synthetic; corporate workdatawas notcopiedtoDEV.",("S1","S2","S3","S4","S6","E2"))
    return p


def main():
    p=build()
    d=b.new_deck("Implemented ISV security diagrams v2")
    implementation(b.security_slide(d))
    identity(d)
    process(d)
    outputs=[b.save(p,"Fabric_ISV_Identity_and_Analytics_v2.pptx"),
             b.save(d,"ISV_Security_Diagrams_v2.pptx")]
    shutil.copy2(b.STAGE/"implementation_v2.air.json",b.OUTPUT/"implementation_v2.air.json")
    (b.OUTPUT/"build_manifest_v2.json").write_text(json.dumps({
        "outputs":outputs,"purpose":"Actual educational implementation, not proposedcustomerarchitecture",
        "sources":b.SOURCES,"classification":"Allslidesownauthored/publicsources/syntheticdata",
        "scope":"Selectedlivecases; noEmbedded/login/productionclaim",
        "liveGearUp":"Search succeeded for revision2; no source slidescollated",
        "nativeEditability":"Groupednodesandattachedconnectors",
    },indent=2)+"\n",encoding="utf-8")
    print(json.dumps(outputs,indent=2))


if __name__=="__main__":
    main()
