"""Publish reviewed decks under unused versions, retaining an immutable local copy."""
import hashlib
import json
from pathlib import Path
import re
import shutil
import time

from pptx import Presentation
from build_customer_deck import STAGE, OUTPUT, SECURITY, PARITY, ICON, census, check

DEST = Path.home() / "OneDrive - Microsoft" / "Documents" / "Microsoft Scout" / "IQ-MCP-RLS-Research" / "Customer-Overview"
REPO = Path(__file__).resolve().parents[2] / "docs" / "customer-overview"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from walk(shape.shapes)


def graph(slide, air):
    nodes = {s.name[5:]: s for s in slide.shapes if s.name.startswith("node-")}
    edges = {s.name[5:]: s for s in slide.shapes if s.name.startswith("edge-")}
    if set(nodes) != {n.id for n in air.nodes} or set(edges) != {e.id for e in air.edges}:
        raise ValueError("Native node/edge IDs differ from authored AIR")
    for edge in air.edges:
        shape = edges[edge.id]
        start = shape._element.xpath(".//a:stCxn")
        end = shape._element.xpath(".//a:endCxn")
        if (len(start) != 1 or len(end) != 1 or
            int(start[0].get("id")) != nodes[edge.from_node].shape_id or
            int(end[0].get("id")) != nodes[edge.to_node].shape_id):
            raise ValueError(f"Incorrect native attachment for {edge.id}")
    for n in air.nodes:
        if n.service_id:
            icon = ICON["model" if n.service_id == "fabric.semantic_model" else "powerbi"]
            blobs = [hashlib.sha256(s.image.blob).hexdigest() for s in walk(nodes[n.id].shapes) if s.shape_type == 13]
            if sha(icon) not in blobs:
                raise ValueError(f"Official asset missing from native group {n.id}")


def signature(slide):
    return [(s.name, int(s.shape_type), s.left, s.top, s.width, s.height,
             s.text if s.has_text_frame else "",
             hashlib.sha256(s.image.blob).hexdigest() if s.shape_type == 13 else "")
            for s in walk(slide.shapes)]


def copy_new(src, dest):
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("xb") as out, src.open("rb") as inp:
        shutil.copyfileobj(inp, out)
    time.sleep(1.5)
    if sha(src) != sha(dest):
        raise ValueError(f"Published bytes changed; clean copy is {src}, destination {dest}")
    if src.suffix.lower() == ".pptx" and census(Presentation(src)) != census(Presentation(dest)):
        raise ValueError(f"Published shape/slide census differs at {dest}")


def main():
    overview = OUTPUT / "Fabric_ISV_Identity_and_Analytics_v1.pptx"
    diagrams = OUTPUT / "ISV_Security_Diagrams_v1.pptx"
    review = json.loads((STAGE / "diagram-review-render-3.json").read_text(encoding="utf-8"))
    editing = json.loads((STAGE / "native-editability_v1.json").read_text(encoding="utf-8-sig"))
    if not review["review"]["passed"] or review["receipt"]["pptx_sha256"] != sha(overview):
        raise ValueError("Architecture review does not bind to this final deck")
    if not editing["passed"] or editing["sourceSha256"].lower() != sha(overview):
        raise ValueError("Native editor observations do not bind to this final deck")
    p, d = Presentation(overview), Presentation(diagrams)
    check(p)
    check(d)
    graph(p.slides[2], SECURITY)
    graph(p.slides[6], PARITY)
    graph(d.slides[0], SECURITY)
    graph(d.slides[2], PARITY)
    for source, target in ((2,0),(3,1),(6,2)):
        if signature(p.slides[source]) != signature(d.slides[target]):
            raise ValueError("Standalone diagram differs from reviewed overview slide")
    DEST.mkdir(parents=True, exist_ok=True)
    existing = [int(m.group(1)) for f in DEST.iterdir()
                if (m := re.search(r"_v(\d+)\.",f.name))]
    version = max(existing, default=0) + 1
    immutable = STAGE / "published" / f"v{version}"
    immutable.mkdir(parents=True, exist_ok=False)
    items = [
        (overview, f"Fabric_ISV_Identity_and_Analytics_v{version}.pptx"),
        (diagrams, f"ISV_Security_Diagrams_v{version}.pptx"),
        (STAGE/"render-3"/"slide-03.png", f"Security_Request_Path_v{version}.png"),
        (STAGE/"render-3"/"slide-04.png", f"Identity_Boundaries_v{version}.png"),
        (STAGE/"render-3"/"slide-07.png", f"Embedded_Comparison_v{version}.png"),
    ]
    published = []
    for source,name in items:
        clean = immutable/name
        copy_new(source,clean)
        dest=DEST/name
        copy_new(clean,dest)
        published.append({"path":str(dest),"staged_copy":str(clean),"sha256":sha(clean)})
    for source,name in items:
        copy_new(source,REPO/name)
    for name in ("security_v1.air.json","parity_v1.air.json"):
        copy_new(OUTPUT/name,REPO/name)
    manifest=json.loads((OUTPUT/"build_manifest_v1.json").read_text(encoding="utf-8"))
    manifest["outputs"]=published
    manifest["status"]="published; exact bytes, full shape census and slide sequence match the staged output"
    manifest["gates"]={
        "request_fit":"14-slide customer narrative; 3-diagram editable pack; login/entitlements, RLS and Embedded differences covered",
        "factual":"Public contracts plus observed synthetic RLS evidence; live IQ and Embedded parity explicitly not claimed",
        "visual":"All 14 slides inspected; final architecture review passed with scoped gpt-6-astra and exact image hashes",
        "native_graph":"All node IDs, edge directions and connection IDs match the small AIRs; grouped icon bytes match official assets",
        "native_editability":"Actual PowerPoint grouped node movement, label edits and automatic attachment passed on both architecture slides",
        "legacy_qa_limitation":"Legacy heuristic did not recurse into grouped icons/text and used a missing catalog path; not treated as an accurate native-group structural verdict",
        "publication":"Exact byte hashes and total recursive shape census/slide sequence matched",
        "fingerprint":"Final published-artifact scan follows publication"
    }
    manifest["classification_by_slide"]={
        "Private Preview":[],"NDA Only":[],"Confidential/PROTECTED":[],
        "customer_shareable_source_material":"Public Microsoft documentation; does not certify output sharing eligibility",
        "our_own":{"overview":list(range(1,15)),"diagram_pack":[1,2,3]}
    }
    local_manifest=immutable/f"build_manifest_v{version}.json"
    local_manifest.write_text(json.dumps(manifest,indent=2)+"\n",encoding="utf-8")
    copy_new(local_manifest,DEST/local_manifest.name)
    print(json.dumps({"version":version,"folder":str(DEST),"published":published,
                      "repo_artifacts":str(REPO)},indent=2))


if __name__ == "__main__":
    main()
